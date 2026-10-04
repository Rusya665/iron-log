import os
import shutil
import sys
import tempfile
import types
import unittest
from unittest.mock import MagicMock

import openpyxl

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from core.models import Exercise, Log
from core.profile_manager import Profile
from core.xlsx_generator import TrainingLogProcessor
from ui.bridge import WebViewBridgeApi


class TestXlsxGenerator(unittest.TestCase):
    """Unit test suite for Excel log generation pipeline and personal records processing."""

    def setUp(self) -> None:
        """
        Create temporary directory and mock data for testing xlsx generation.

        :return: None
        """
        self.temp_dir = tempfile.mkdtemp()
        self.output_xlsx = os.path.join(self.temp_dir, "test_training_log.xlsx")

        self.exercises = [
            Exercise(id="bench_press", display_name="Bench Press"),
            Exercise(id="squat", display_name="Squat"),
            Exercise(id="pullups", display_name="Pull-ups"),
        ]

        self.bodymass_log = {
            "2026-01-01": 80.0,
            "2026-01-03": 80.2,
            "2026-01-05": 79.8,
        }

        self.user_data = {
            "2026-01-01": {
                "day": 1,
                "bench_press": Log(reps=[10, 8, 6], mass=[80.0, 90.0, 100.0]),
                "pullups": Log(reps=[12, 10], mass=[0.0, 0.0]),
            },
            "2026-01-03": {
                "day": "PR",
                "bench_press": Log(reps=[1], mass=[110.0]),
                "squat": Log(reps=[1], mass=[145.0]),
            },
            "2026-01-05": {
                "day": 2,
                "squat": Log(reps=[5, 5, 5], mass=[120.0, 125.0, 130.0]),
            },
        }

        self.user_profile = {
            "name": "Test Athlete",
            "sex": "male",
            "age": 28,
            "mass": 80.0,
            "show_pr": True,
            "show_standards": True,
            "show_milestones": True,
        }

    def tearDown(self) -> None:
        """
        Clean up temporary directory and files.

        :return: None
        """
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_write_personal_records_no_name_error(self) -> None:
        """
        Verify write_personal_records executes without NameError: name 'Tuple' is not defined.

        Regression test for missing Tuple import in core.xlsx_generator.
        :return: None
        """
        processor = TrainingLogProcessor(
            self.output_xlsx,
            self.exercises,
            self.user_data,
            self.bodymass_log,
            self.user_profile,
        )
        # Calling write_personal_records previously raised NameError: name 'Tuple' is not defined
        try:
            processor.write_personal_records()
        except NameError as e:
            self.fail(f"write_personal_records raised NameError: {e}")

    def test_generate_xlsx_full_pipeline(self) -> None:
        """
        Verify complete workbook generation pipeline executes end-to-end and outputs valid xlsx.

        :return: None
        """
        processor = TrainingLogProcessor(
            self.output_xlsx,
            self.exercises,
            self.user_data,
            self.bodymass_log,
            self.user_profile,
        )
        processor.validate_data()
        processor.write_headers()
        processor.process_data(self.user_data)
        processor.write_calculations()
        processor.generate_charts()
        processor.write_definitions()
        processor.write_personal_records()
        processor.write_user_profile()
        processor.save()

        self.assertTrue(os.path.exists(self.output_xlsx))
        self.assertGreater(os.path.getsize(self.output_xlsx), 0)

        # Validate sheet structure and contents via openpyxl
        wb = openpyxl.load_workbook(self.output_xlsx)
        expected_sheets = [
            "Progress_Charts",
            "Data_Log",
            "Definitions",
            "Calculations",
            "User_Profile",
            "Personal_Records",
        ]
        for sheet_name in expected_sheets:
            self.assertIn(sheet_name, wb.sheetnames)

        # Check Personal_Records sheet headers
        ws_pr = wb["Personal_Records"]
        self.assertEqual(ws_pr.cell(row=1, column=1).value, "Training PRs")
        self.assertEqual(ws_pr.cell(row=1, column=2).value, "PR Mass (kg)")
        self.assertEqual(ws_pr.cell(row=1, column=3).value, "Date Achieved")
        self.assertEqual(ws_pr.cell(row=1, column=4).value, "Strength Level")

        self.assertEqual(ws_pr.cell(row=1, column=6).value, "PR Session Bests")
        self.assertEqual(ws_pr.cell(row=1, column=7).value, "PR Mass (kg)")
        self.assertEqual(ws_pr.cell(row=1, column=8).value, "Date Achieved")
        self.assertEqual(ws_pr.cell(row=1, column=9).value, "Strength Level")

        # Check Personal_Records values
        # Bench Press: 100 kg on training day 2026-01-01, 110 kg on PR session 2026-01-03
        self.assertEqual(ws_pr.cell(row=2, column=1).value, "Bench Press")
        self.assertEqual(ws_pr.cell(row=2, column=2).value, 100)
        self.assertEqual(ws_pr.cell(row=2, column=3).value, "2026-01-01")

        self.assertEqual(ws_pr.cell(row=2, column=6).value, "Bench Press")
        self.assertEqual(ws_pr.cell(row=2, column=7).value, 110)
        self.assertEqual(ws_pr.cell(row=2, column=8).value, "2026-01-03")

        # Squat: 130 kg on training day 2026-01-05, 145 kg on PR session 2026-01-03
        self.assertEqual(ws_pr.cell(row=3, column=1).value, "Squat")
        self.assertEqual(ws_pr.cell(row=3, column=2).value, 130)
        self.assertEqual(ws_pr.cell(row=3, column=3).value, "2026-01-05")

        self.assertEqual(ws_pr.cell(row=3, column=6).value, "Squat")
        self.assertEqual(ws_pr.cell(row=3, column=7).value, 145)
        self.assertEqual(ws_pr.cell(row=3, column=8).value, "2026-01-03")

        # Check User_Profile sheet content
        ws_prof = wb["User_Profile"]
        self.assertEqual(ws_prof.cell(row=2, column=1).value, "Name")
        self.assertEqual(ws_prof.cell(row=2, column=2).value, "Test Athlete")
        self.assertEqual(ws_prof.cell(row=4, column=1).value, "Sex")
        self.assertEqual(ws_prof.cell(row=4, column=2).value, "male")

        wb.close()

    def test_bridge_generate_excel_success(self) -> None:
        """
        Verify WebViewBridgeApi.generate_excel runs generation without raising NameError or failing.

        :return: None
        """
        p = Profile(
            name="Bridge Athlete",
            sessions_dir=self.temp_dir,
            output_dir=self.temp_dir,
            sex="male",
            mass=80.0,
        )

        mock_sess = types.SimpleNamespace(
            EXERCISE_REGISTRY=self.exercises,
            USER_DATA=self.user_data,
            BODYMASS_LOG=self.bodymass_log,
        )

        api = WebViewBridgeApi()
        api.manager.get_active_profile = MagicMock(return_value=p)
        api._load_sessions = MagicMock(return_value=(mock_sess, "dummy_sessions.py"))

        res = api.generate_excel()
        self.assertTrue(res.get("success"), f"Bridge generate_excel failed: {res}")
        self.assertIn("time", res)

        generated_files = [f for f in os.listdir(self.temp_dir) if f.endswith(".xlsx")]
        self.assertGreater(len(generated_files), 0)

    def test_validate_data_mismatched_reps_mass(self) -> None:
        """
        Verify validate_data raises ValueError when rep and mass counts do not match.

        :return: None
        """
        bad_data = {
            "2026-01-01": {
                "day": 1,
                "bench_press": Log(reps=[10, 8], mass=[80.0, 90.0, 100.0]),
            }
        }
        processor = TrainingLogProcessor(
            self.output_xlsx,
            self.exercises,
            bad_data,
            self.bodymass_log,
            self.user_profile,
        )
        with self.assertRaises(ValueError):
            processor.validate_data()


if __name__ == "__main__":
    unittest.main()
